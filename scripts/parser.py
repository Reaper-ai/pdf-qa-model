import os
import json
import pymupdf
import easyocr
from PIL import Image
from docx import Document
from pptx import Presentation
from typing import List, Dict, Any
import numpy as np

class UniversalParser:
    def __init__(self):
        self.ocr_reader = easyocr.Reader(['en'], gpu=False)

    def parse(self, file_path: str) -> List[Dict[str, Any]]:
        """
        Takes a file path, extracts the filename, and routes to the correct parser.
        """
        # Automatically extract the file name from the path
        file_name = os.path.basename(file_path)
        
        ext = file_name.split('.')[-1].lower() if '.' in file_name else ''
        
        if ext == 'pdf':
            return self._parse_pdf(file_path, file_name)
        elif ext in ['docx', 'doc']:
            return self._parse_docx(file_path, file_name)
        elif ext in ['pptx', 'ppt']:
            return self._parse_pptx(file_path, file_name)
        elif ext in ['jpg', 'jpeg', 'png']:
            return self._parse_image(file_path, file_name)
        elif ext == 'json':
            return self._parse_json(file_path, file_name)
        elif ext in ['txt', 'md', '']: 
            return self._parse_txt(file_path, file_name)
        else:
            raise ValueError(f"Unsupported file format: {ext}")

    def _parse_pdf(self, file_path: str, file_name: str) -> List[Dict[str, Any]]:
        parsed_data = []
        doc = pymupdf.open(file_path)
        
        # Extract native PDF metadata
        native_meta = doc.metadata if doc.metadata else {}
        author = native_meta.get("author", "Unknown")
        title = native_meta.get("title", "Unknown")
        creation_date = native_meta.get("creationDate", "Unknown")

        for page_num, page in enumerate(doc, start=1):
            text = page.get_text("text")
            parsed_data.append({
                "text": text,
                "metadata": {
                    "source": file_name, 
                    "page_number": page_num, 
                    "type": "pdf",
                    "author": author,
                    "title": title,
                    "creation_date": creation_date
                }
            })
        return parsed_data

    def _parse_docx(self, file_path: str, file_name: str) -> List[Dict[str, Any]]:
        doc = Document(file_path)
        
        # Extract native DOCX metadata
        core_props = doc.core_properties
        author = core_props.author or "Unknown"
        title = core_props.title or "Unknown"
        created = core_props.created.strftime("%Y-%m-%d") if core_props.created else "Unknown"

        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        parsed_data = []
        block_size = 5
        
        for i in range(0, len(paragraphs), block_size):
            text_block = "\n".join(paragraphs[i:i + block_size])
            parsed_data.append({
                "text": text_block,
                "metadata": {
                    "source": file_name, 
                    "block": (i // block_size) + 1, 
                    "type": "docx",
                    "author": author,
                    "title": title,
                    "creation_date": created
                }
            })
        return parsed_data

    def _parse_pptx(self, file_path: str, file_name: str) -> List[Dict[str, Any]]:
        prs = Presentation(file_path)
        
        # Extract native PPTX metadata
        core_props = prs.core_properties
        author = core_props.author or "Unknown"
        title = core_props.title or "Unknown"
        created = core_props.created.strftime("%Y-%m-%d") if core_props.created else "Unknown"

        parsed_data = []
        for slide_num, slide in enumerate(prs.slides, start=1):
            text_runs = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text_runs.append(shape.text)
            
            parsed_data.append({
                "text": "\n".join(text_runs),
                "metadata": {
                    "source": file_name, 
                    "slide_number": slide_num, 
                    "type": "pptx",
                    "author": author,
                    "title": title,
                    "creation_date": created
                }
            })
        return parsed_data
    
    def _parse_image(self, file_path: str, file_name: str) -> List[Dict[str, Any]]:
        img = Image.open(file_path).convert("RGB")
        
        results = self.ocr_reader.readtext(np.array(img), detail=0)
        text = " ".join(results)
        
        return [{
            "text": text,
            "metadata": {"source": file_name, "type": "image"}
        }]

    def _parse_json(self, file_path: str, file_name: str) -> List[Dict[str, Any]]:
        with open(file_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        
        text_dump = json.dumps(data, indent=2)
        return [{
            "text": text_dump,
            "metadata": {"source": file_name, "type": "json"}
        }]

    
    def _parse_txt(self, file_path: str, file_name: str) -> List[Dict[str, Any]]:
        with open(file_path, mode='r', encoding='utf-8', errors='ignore') as f:
            text_content = f.read()

        paragraphs = [p.strip() for p in text_content.split('\n\n') if p.strip()]
        
        if not paragraphs and text_content.strip():
            paragraphs = [text_content.strip()]

        parsed_data = []
        block_size = 10 
        
        for i in range(0, len(paragraphs), block_size):
            text_block = "\n\n".join(paragraphs[i:i + block_size])
            
            parsed_data.append({
                "text": text_block,
                "metadata": {
                    "source": file_name, 
                    "block": (i // block_size) + 1, 
                    "type": "txt"
                }
            })
            
        return parsed_data